import os
import sys
import yaml
from docx import Document
from docx.shared import Inches

def pedir_dia_y_reemplazar_fecha(doc):
    """Busca el campo de fecha en la cabecera y actualiza el día."""
    for p in doc.paragraphs:
        if "Fecha:" in p.text and "/20" in p.text:
            partes = p.text.split("Fecha:")
            if len(partes) > 1:
                fecha_actual = partes[1].strip()
                print(f"📅 Fecha encontrada: {fecha_actual}")
                dia = input("👉 Introduce el día que quieres poner (número): ").zfill(2)
                nueva_fecha = f"{dia}{fecha_actual}"
                p.text = f"{partes[0]}Fecha: {nueva_fecha}"

def rellenar_registro(archivo_entrada, config_path="config.yaml"):
    # Verificaciones
    if not os.path.exists(archivo_entrada):
        print(f"❌ El archivo {archivo_entrada} no existe.")
        return
    if not os.path.exists(config_path):
        print(f"❌ El archivo de configuración {config_path} no existe.")
        return

    # Cargar configuración YAML
    with open(config_path, "r", encoding="utf-8") as f:
        config = yaml.safe_load(f)

    nombre_apellidos = config.get("nombre_apellidos", "")
    nif = config.get("nif", "")
    naf = config.get("naf", "")
    hora_entrada_mañana = config.get("hora_entrada_mañana", "9:00")
    hora_salida_mañana = config.get("hora_salida_mañana", "")
    hora_entrada_tarde = config.get("hora_entrada_tarde", "")
    hora_salida_tarde = config.get("hora_salida_tarde", "17:00")
    firma_path = config.get("firma", "sign/sign.png")

    if not os.path.exists(firma_path):
        print(f"❌ La firma {firma_path} no existe.")
        return

    # Cargar documento
    doc = Document(archivo_entrada)

    # 1️⃣ Rellenar nombre, NIF y NAF
    for p in doc.paragraphs:
        if "Nombre y apellidos:" in p.text:
            p.text = f"Nombre y apellidos: {nombre_apellidos}"
        if "NIF:" in p.text or "NAF:" in p.text:
            p.text = f"NIF: {nif}                   NAF: {naf}"

    # 2️⃣ Actualizar fecha en cabecera
    pedir_dia_y_reemplazar_fecha(doc)

    # 3️⃣ Localizar la tabla principal (la que tiene más filas)
    tabla_principal = max(doc.tables, key=lambda t: len(t.rows))

    # 4️⃣ Rellenar horas y firma diaria
    for row in tabla_principal.rows:
        primera_celda = row.cells[0].text.strip()
        if not primera_celda.isdigit():  # solo las filas con número de día
            continue
        if len(row.cells) >= 6:
            # Entrada mañana
            if row.cells[1].text.strip() == "" and hora_entrada_mañana:
                row.cells[1].text = hora_entrada_mañana
            # Salida mañana
            row.cells[2].text = hora_salida_mañana if hora_salida_mañana else ""
            # Entrada tarde
            row.cells[3].text = hora_entrada_tarde if hora_entrada_tarde else ""
            # Salida tarde
            if row.cells[4].text.strip() == "" and hora_salida_tarde:
                row.cells[4].text = hora_salida_tarde
            # Firma diaria
            firma_cell = row.cells[5]
            firma_cell.text = ""
            run = firma_cell.paragraphs[0].add_run()
            run.add_picture(firma_path, width=Inches(1))

    # 5️⃣ Guardar documento final
    os.makedirs("output", exist_ok=True)
    salida = os.path.join("output", os.path.basename(archivo_entrada))
    doc.save(salida)
    print(f"✅ Archivo generado correctamente: {salida}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python rellenar_registro.py <archivo_entrada.docx> [config.yaml]")
    else:
        archivo = sys.argv[1]
        config = sys.argv[2] if len(sys.argv) > 2 else "config.yaml"
        rellenar_registro(archivo, config)
